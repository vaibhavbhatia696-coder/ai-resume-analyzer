import re
import html
from pathlib import Path

import streamlit as st


# Optional libraries
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except ImportError:
    TfidfVectorizer = None
    cosine_similarity = None


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    .main {
        padding-top: 1rem;
    }

    .title {
        font-size: 42px;
        font-weight: 800;
        color: #2563eb;
        margin-bottom: 0;
    }

    .subtitle {
        color: #64748b;
        font-size: 17px;
        margin-bottom: 25px;
    }

    .metric-card {
        background: linear-gradient(135deg, #eff6ff, #dbeafe);
        border-radius: 15px;
        padding: 20px;
        text-align: center;
        border: 1px solid #bfdbfe;
    }

    .metric-number {
        font-size: 34px;
        font-weight: 800;
        color: #1d4ed8;
    }

    .metric-label {
        font-size: 14px;
        color: #475569;
        font-weight: 600;
    }

    .score-card {
        background: linear-gradient(135deg, #eff6ff, #dbeafe);
        border-radius: 18px;
        padding: 28px;
        text-align: center;
        border: 1px solid #bfdbfe;
        min-height: 190px;
    }

    .score-number {
        font-size: 55px;
        font-weight: 800;
        color: #2563eb;
        line-height: 1;
        margin: 15px 0;
    }

    .score-label {
        color: #64748b;
        font-size: 14px;
        font-weight: 600;
    }

    .skill {
        display: inline-block;
        background: #dbeafe;
        color: #1e40af;
        border-radius: 999px;
        padding: 7px 12px;
        margin: 4px;
        font-size: 13px;
        font-weight: 600;
    }

    .matched {
        display: inline-block;
        background: #dcfce7;
        color: #166534;
        border-radius: 999px;
        padding: 7px 12px;
        margin: 4px;
        font-size: 13px;
        font-weight: 600;
    }

    .missing {
        display: inline-block;
        background: #fee2e2;
        color: #991b1b;
        border-radius: 999px;
        padding: 7px 12px;
        margin: 4px;
        font-size: 13px;
        font-weight: 600;
    }

    .suggestion {
        background: #f8fafc;
        border-left: 4px solid #2563eb;
        padding: 12px 15px;
        margin: 8px 0;
        border-radius: 7px;
        color: #334155;
    }

    .text-box {
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        border-radius: 12px;
        padding: 18px;
        max-height: 450px;
        overflow-y: auto;
        white-space: pre-wrap;
        font-size: 14px;
        line-height: 1.6;
    }

    .info-box {
        padding: 15px;
        background: #eff6ff;
        border-radius: 10px;
        border: 1px solid #bfdbfe;
        color: #1e3a8a;
    }

    .ats-item {
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        border-radius: 10px;
        padding: 12px 15px;
        margin-bottom: 8px;
        font-weight: 600;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SKILLS DATABASE
# ============================================================

SKILLS = {

    # Programming
    "python": ["python"],
    "java": ["java"],
    "javascript": ["javascript", "js"],
    "typescript": ["typescript", "ts"],
    "c": ["c programming"],
    "c++": ["c++", "cpp"],
    "c#": ["c#", "c sharp"],
    "php": ["php"],
    "ruby": ["ruby"],
    "go": ["golang", "go language"],
    "rust": ["rust"],
    "kotlin": ["kotlin"],
    "swift": ["swift"],

    # Web
    "html": ["html", "html5"],
    "css": ["css", "css3"],
    "react": ["react", "reactjs", "react.js"],
    "angular": ["angular", "angularjs"],
    "vue": ["vue", "vuejs"],
    "node.js": ["node.js", "nodejs", "node js"],
    "express.js": ["express.js", "expressjs", "express js"],
    "next.js": ["next.js", "nextjs", "next js"],

    # Data / AI
    "machine learning": ["machine learning", "machine-learning"],
    "deep learning": ["deep learning", "deep-learning"],
    "artificial intelligence": [
        "artificial intelligence",
        "ai",
    ],
    "data science": ["data science", "data scientist"],
    "data analysis": [
        "data analysis",
        "data analytics",
        "data analyst",
    ],
    "nlp": [
        "nlp",
        "natural language processing",
    ],
    "computer vision": ["computer vision"],
    "tensorflow": ["tensorflow"],
    "pytorch": ["pytorch"],
    "scikit-learn": [
        "scikit-learn",
        "sklearn",
    ],
    "pandas": ["pandas"],
    "numpy": ["numpy"],
    "opencv": ["opencv"],
    "matplotlib": ["matplotlib"],

    # Databases
    "sql": ["sql"],
    "mysql": ["mysql"],
    "postgresql": [
        "postgresql",
        "postgres",
    ],
    "mongodb": [
        "mongodb",
        "mongo db",
        "mongo",
    ],
    "sqlite": ["sqlite"],
    "redis": ["redis"],
    "oracle": [
        "oracle database",
        "oracle db",
    ],

    # Cloud / DevOps
    "aws": [
        "aws",
        "amazon web services",
    ],
    "azure": [
        "azure",
        "microsoft azure",
    ],
    "google cloud": [
        "google cloud",
        "gcp",
    ],
    "docker": ["docker"],
    "kubernetes": [
        "kubernetes",
        "k8s",
    ],
    "jenkins": ["jenkins"],
    "git": ["git"],
    "github": ["github"],
    "gitlab": ["gitlab"],
    "ci/cd": [
        "ci/cd",
        "cicd",
        "continuous integration",
    ],
    "terraform": ["terraform"],

    # Software Engineering
    "rest api": [
        "rest api",
        "restful api",
        "restful",
    ],
    "api development": ["api development"],
    "object oriented programming": [
        "object oriented programming",
        "oop",
    ],
    "data structures": [
        "data structures",
        "data structure",
    ],
    "algorithms": [
        "algorithms",
        "algorithm",
    ],
    "unit testing": [
        "unit testing",
        "unit test",
    ],
    "agile": ["agile"],
    "scrum": ["scrum"],
    "software development": [
        "software development",
    ],

    # Other
    "excel": [
        "excel",
        "microsoft excel",
    ],
    "power bi": [
        "power bi",
        "powerbi",
    ],
    "tableau": ["tableau"],
    "communication": [
        "communication skills",
        "communication",
    ],
    "leadership": [
        "leadership",
        "leadership skills",
    ],
    "problem solving": [
        "problem solving",
        "problem-solving",
    ],
    "project management": [
        "project management",
    ],
}


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text: str) -> str:

    if not text:
        return ""

    text = text.lower()

    text = re.sub(
        r"[/|,;:()\[\]{}]",
        " ",
        text,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(uploaded_file) -> str:

    if PdfReader is None:
        raise RuntimeError(
            "PDF support unavailable. Run: pip install pypdf"
        )

    try:

        uploaded_file.seek(0)

        reader = PdfReader(uploaded_file)

        pages = []

        for page in reader.pages:

            try:

                page_text = page.extract_text() or ""

                pages.append(page_text)

            except Exception:
                continue

        text = "\n\n".join(pages).strip()

        if not text:
            raise ValueError(
                "PDF se readable text nahi mila. "
                "Agar PDF scanned/image-based hai, OCR ki zarurat hogi."
            )

        return text

    except Exception as e:

        raise RuntimeError(
            f"PDF read karne mein problem: {e}"
        )


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(uploaded_file) -> str:

    if Document is None:
        raise RuntimeError(
            "DOCX support unavailable. Run: "
            "pip install python-docx"
        )

    try:

        uploaded_file.seek(0)

        document = Document(uploaded_file)

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    paragraphs.append(
                        " | ".join(row_text)
                    )

        text = "\n".join(paragraphs).strip()

        if not text:
            raise ValueError(
                "DOCX file mein readable text nahi mila."
            )

        return text

    except Exception as e:

        raise RuntimeError(
            f"DOCX read karne mein problem: {e}"
        )


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt_text(uploaded_file) -> str:

    try:

        uploaded_file.seek(0)

        raw = uploaded_file.read()

        if isinstance(raw, bytes):

            for encoding in [
                "utf-8",
                "utf-8-sig",
                "latin-1",
            ]:

                try:

                    text = raw.decode(encoding)

                    break

                except UnicodeDecodeError:
                    continue

            else:

                raise ValueError(
                    "TXT file ki encoding read nahi ho saki."
                )

        else:

            text = str(raw)

        text = text.strip()

        if not text:
            raise ValueError(
                "TXT file empty hai."
            )

        return text

    except Exception as e:

        raise RuntimeError(
            f"TXT read karne mein problem: {e}"
        )


# ============================================================
# UNIVERSAL FILE EXTRACTION
# ============================================================

def extract_resume_text(uploaded_file) -> str:

    filename = uploaded_file.name.lower()

    extension = Path(filename).suffix

    if extension == ".pdf":
        return extract_pdf_text(uploaded_file)

    if extension == ".docx":
        return extract_docx_text(uploaded_file)

    if extension == ".txt":
        return extract_txt_text(uploaded_file)

    raise ValueError(
        "Unsupported file format. "
        "Sirf PDF, DOCX aur TXT files allowed hain."
    )


# ============================================================
# SKILL DETECTION
# ============================================================

def contains_skill(text: str, skill: str) -> bool:

    normalized = normalize_text(text)

    aliases = SKILLS.get(
        skill,
        [skill]
    )

    for alias in aliases:

        alias = normalize_text(alias)

        if not alias:
            continue

        pattern = (
            rf"(?<![a-z0-9])"
            rf"{re.escape(alias)}"
            rf"(?![a-z0-9])"
        )

        if re.search(
            pattern,
            normalized
        ):
            return True

    return False


def detect_skills(text: str):

    detected = []

    for skill in SKILLS:

        if contains_skill(
            text,
            skill
        ):

            detected.append(skill)

    return sorted(detected)


# ============================================================
# JOB DESCRIPTION SKILL EXTRACTION
# ============================================================

def extract_required_skills(
    job_description: str
):

    if not job_description:
        return []

    return detect_skills(
        job_description
    )


# ============================================================
# TEXT SIMILARITY
# ============================================================

def calculate_text_similarity(
    resume_text: str,
    job_description: str
) -> float:

    if (
        not resume_text.strip()
        or not job_description.strip()
    ):
        return 0.0

    if (
        TfidfVectorizer is None
        or cosine_similarity is None
    ):
        return 0.0

    try:

        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            max_features=10000,
        )

        matrix = vectorizer.fit_transform(
            [
                resume_text,
                job_description
            ]
        )

        similarity = cosine_similarity(
            matrix[0:1],
            matrix[1:2]
        )[0][0]

        return max(
            0.0,
            min(
                100.0,
                similarity * 100
            )
        )

    except Exception:
        return 0.0


# ============================================================
# SKILL MATCH SCORE
# ============================================================

def calculate_skill_score(
    resume_skills,
    required_skills
) -> float:

    if not required_skills:
        return 0.0

    resume_set = set(
        resume_skills
    )

    required_set = set(
        required_skills
    )

    matched = resume_set.intersection(
        required_set
    )

    return (
        len(matched)
        / len(required_set)
    ) * 100


# ============================================================
# FINAL MATCH SCORE
# ============================================================

def calculate_final_score(
    resume_text,
    job_description,
    resume_skills,
    required_skills
) -> float:

    text_score = calculate_text_similarity(
        resume_text,
        job_description
    )

    skill_score = calculate_skill_score(
        resume_skills,
        required_skills
    )

    if required_skills:

        final_score = (
            skill_score * 0.60
            +
            text_score * 0.40
        )

    else:

        final_score = text_score

    return round(
        max(
            0.0,
            min(
                100.0,
                final_score
            )
        ),
        1
    )


# ============================================================
# ATS RESUME ANALYSIS
# ============================================================

def analyze_ats_resume(
    resume_text
):

    text = normalize_text(
        resume_text
    )

    checks = {

        "Contact Information":
            bool(
                re.search(
                    r"\b[\w\.-]+@[\w\.-]+\.\w+\b",
                    text
                )
            ),

        "Skills":
            "skills" in text,

        "Education":
            (
                "education" in text
                or "b.tech" in text
                or "bachelor" in text
                or "degree" in text
            ),

        "Projects":
            (
                "projects" in text
                or "project" in text
            ),

        "Experience":
            (
                "experience" in text
                or "internship" in text
                or "intern" in text
            ),

        "Certifications":
            (
                "certification" in text
                or "certificate" in text
            ),

        "Achievements":
            (
                "achievement" in text
                or "award" in text
            ),

    }

    # Resume length
    word_count = len(
        re.findall(
            r"\b\w+\b",
            resume_text
        )
    )

    checks["Readable Length"] = (
        250 <= word_count <= 1200
    )

    # Technical skills
    checks["Technical Skills"] = (
        len(
            detect_skills(
                resume_text
            )
        ) >= 3
    )

    passed = sum(
        1
        for value in checks.values()
        if value
    )

    total = len(checks)

    score = round(
        (
            passed
            /
            total
        ) * 100
    )

    return score, checks


# ============================================================
# ATS KEYWORD ANALYSIS
# ============================================================

def calculate_keyword_coverage(
    resume_text,
    job_description
):

    if not job_description.strip():
        return 0.0, [], []

    required_skills = extract_required_skills(
        job_description
    )

    resume_skills = detect_skills(
        resume_text
    )

    if not required_skills:
        return 0.0, [], []

    matched = sorted(
        set(required_skills)
        &
        set(resume_skills)
    )

    missing = sorted(
        set(required_skills)
        -
        set(resume_skills)
    )

    coverage = (
        len(matched)
        /
        len(required_skills)
    ) * 100

    return (
        round(coverage, 1),
        matched,
        missing
    )


# ============================================================
# SUGGESTIONS
# ============================================================

def generate_suggestions(
    resume_text,
    job_description,
    resume_skills,
    required_skills,
    missing_skills
):

    suggestions = []

    normalized_resume = normalize_text(
        resume_text
    )

    if not resume_text.strip():

        suggestions.append(
            "Resume text empty hai. "
            "Ek readable PDF, DOCX ya TXT upload karein."
        )

        return suggestions

    if job_description.strip():

        if not required_skills:

            suggestions.append(
                "Job description mein predefined "
                "technical skills detect nahi hui. "
                "Important keywords manually review karein."
            )

        if missing_skills:

            top_missing = missing_skills[:8]

            suggestions.append(
                "Agar ye skills aapke actual experience "
                "mein hain, toh unhe resume ke Skills "
                "aur Experience sections mein clear "
                "keywords ke saath mention karein: "
                + ", ".join(top_missing)
                + "."
            )

        if "experience" not in normalized_resume:

            suggestions.append(
                "Resume mein Experience / Projects section "
                "clearly add karein aur responsibilities ke "
                "saath measurable results mention karein."
            )

        if "project" not in normalized_resume:

            suggestions.append(
                "Relevant projects add karein, especially "
                "projects jo target job description se "
                "directly related hon."
            )

        if "education" not in normalized_resume:

            suggestions.append(
                "Education section ko clearly structure karein."
            )

        if "skills" not in normalized_resume:

            suggestions.append(
                "Dedicated Skills section maintain karein."
            )

    else:

        suggestions.append(
            "Accurate job-match analysis ke liye "
            "Job Description paste karein."
        )

    unique = []

    for suggestion in suggestions:

        if suggestion not in unique:
            unique.append(suggestion)

    return unique


# ============================================================
# HTML HELPERS
# ============================================================

def render_skills(
    skills,
    css_class="skill"
):

    if not skills:

        return (
            "<span style='color:#64748b;'>"
            "None detected"
            "</span>"
        )

    output = ""

    for skill in skills:

        safe_skill = html.escape(
            skill
        )

        output += (
            f'<span class="{css_class}">'
            f'{safe_skill}'
            f'</span>'
        )

    return output


# ============================================================
# SESSION STATE
# ============================================================

if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""

if "resume_name" not in st.session_state:
    st.session_state.resume_name = ""


# ============================================================
# HEADER
# ============================================================

st.markdown(
    '<div class="title">'
    '📄 AI Resume Analyzer'
    '</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="subtitle">'
    'Upload your resume and compare it with a job description.'
    '</div>',
    unsafe_allow_html=True,
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header(
        "⚙️ Resume Upload"
    )

    uploaded_file = st.file_uploader(
        "Upload Resume",
        type=[
            "pdf",
            "docx",
            "txt"
        ],
        help=(
            "PDF, DOCX aur TXT files supported hain."
        ),
    )

    if uploaded_file is not None:

        if st.button(
            "📥 Extract Resume",
            use_container_width=True,
        ):

            with st.spinner(
                "Resume read ho raha hai..."
            ):

                try:

                    extracted = extract_resume_text(
                        uploaded_file
                    )

                    st.session_state.resume_text = (
                        extracted
                    )

                    st.session_state.resume_name = (
                        uploaded_file.name
                    )

                    st.success(
                        "Resume successfully loaded: "
                        f"{uploaded_file.name}"
                    )

                except Exception as e:

                    st.error(
                        str(e)
                    )

    st.divider()

    if st.session_state.resume_text:

        st.success(
            "✅ Resume loaded"
        )

        if st.session_state.resume_name:

            st.caption(
                st.session_state.resume_name
            )

        if st.button(
            "🗑️ Clear Resume",
            use_container_width=True,
        ):

            st.session_state.resume_text = ""
            st.session_state.resume_name = ""

            st.rerun()

    st.divider()

    st.markdown(
        """
        **Supported formats**

        📄 PDF

        📝 DOCX

        📃 TXT

        **No API key required.**
        """
    )


# ============================================================
# MAIN INPUT AREA
# ============================================================

left_col, right_col = st.columns(2)


with left_col:

    st.subheader(
        "📄 Resume"
    )

    if st.session_state.resume_text:

        resume_preview = (
            st.session_state.resume_text
        )

        st.markdown(
            f'<div class="text-box">'
            f'{html.escape(resume_preview)}'
            f'</div>',
            unsafe_allow_html=True,
        )

    else:

        st.info(
            "Sidebar se PDF, DOCX ya TXT resume "
            "upload karke 'Extract Resume' par click karein."
        )


with right_col:

    st.subheader(
        "💼 Job Description"
    )

    job_description = st.text_area(
        "Paste the complete job description here",
        height=450,
        placeholder=(
            "Example:\n\n"
            "We are looking for a Python Developer "
            "with experience in Django, REST APIs, "
            "SQL, Git and AWS..."
        ),
        label_visibility="collapsed",
    )


# ============================================================
# ANALYZE BUTTON
# ============================================================

st.divider()

analyze_col1, analyze_col2, analyze_col3 = st.columns(
    [1, 2, 1]
)


with analyze_col2:

    analyze_clicked = st.button(
        "🔍 Analyze Resume",
        type="primary",
        use_container_width=True,
    )


# ============================================================
# ANALYSIS
# ============================================================

if analyze_clicked:

    resume_text = (
        st.session_state.resume_text.strip()
    )

    job_description = (
        job_description.strip()
    )

    if not resume_text:

        st.error(
            "Pehle resume upload karke extract karein."
        )

        st.stop()

    if not job_description:

        st.error(
            "Job Description paste karna zaroori hai."
        )

        st.stop()

    with st.spinner(
        "Resume analyze ho raha hai..."
    ):

        # ----------------------------------------------------
        # Skills
        # ----------------------------------------------------

        resume_skills = detect_skills(
            resume_text
        )

        required_skills = extract_required_skills(
            job_description
        )

        matched_skills = sorted(
            set(resume_skills)
            &
            set(required_skills)
        )

        missing_skills = sorted(
            set(required_skills)
            -
            set(resume_skills)
        )

        # ----------------------------------------------------
        # Scores
        # ----------------------------------------------------

        text_score = calculate_text_similarity(
            resume_text,
            job_description
        )

        skill_score = calculate_skill_score(
            resume_skills,
            required_skills
        )

        final_score = calculate_final_score(
            resume_text,
            job_description,
            resume_skills,
            required_skills
        )

        # ----------------------------------------------------
        # ATS
        # ----------------------------------------------------

        ats_score, ats_checks = (
            analyze_ats_resume(
                resume_text
            )
        )

        keyword_coverage, keyword_matched, keyword_missing = (
            calculate_keyword_coverage(
                resume_text,
                job_description
            )
        )

        # ----------------------------------------------------
        # Suggestions
        # ----------------------------------------------------

        suggestions = generate_suggestions(
            resume_text,
            job_description,
            resume_skills,
            required_skills,
            missing_skills
        )


    # ========================================================
    # SCORE
    # ========================================================

    st.subheader(
        "📊 Analysis Result"
    )

    score_col1, score_col2, score_col3 = st.columns(3)


    with score_col1:

        st.markdown(
            f"""
            <div class="metric-card">

                <div class="metric-number">
                    {final_score}%
                </div>

                <div class="metric-label">
                    Overall Match Score
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    with score_col2:

        st.markdown(
            f"""
            <div class="metric-card">

                <div class="metric-number">
                    {skill_score:.1f}%
                </div>

                <div class="metric-label">
                    Skill Match
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    with score_col3:

        st.markdown(
            f"""
            <div class="metric-card">

                <div class="metric-number">
                    {text_score:.1f}%
                </div>

                <div class="metric-label">
                    Text Similarity
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    st.progress(
        int(round(final_score)),
        text=(
            f"Overall Match: {final_score}%"
        ),
    )


    # ========================================================
    # ATS RESUME SCORE
    # ========================================================

    st.divider()

    st.subheader(
        "🤖 ATS Resume Analysis"
    )

    ats_col1, ats_col2 = st.columns(
        [1, 2]
    )


    with ats_col1:

        st.markdown(
            f"""
            <div class="score-card">

                <div style="
                    color:#64748b;
                    font-size:13px;
                    font-weight:700;
                ">
                    ATS FRIENDLINESS
                </div>

                <div class="score-number">
                    {ats_score}%
                </div>

                <div class="score-label">
                    Resume structure analysis
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    with ats_col2:

        st.markdown(
            "#### 📋 Resume Structure"
        )

        for section, status in ats_checks.items():

            if status:

                st.markdown(
                    f"""
                    <div class="ats-item">
                        🟢 {html.escape(section)}
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            else:

                st.markdown(
                    f"""
                    <div class="ats-item">
                        🔴 {html.escape(section)}
                    </div>
                    """,
                    unsafe_allow_html=True,
                )


    # ========================================================
    # KEYWORD COVERAGE
    # ========================================================

    st.divider()

    st.subheader(
        "🔑 ATS Keyword Coverage"
    )

    keyword_col1, keyword_col2, keyword_col3 = st.columns(3)


    with keyword_col1:

        st.markdown(
            f"""
            <div class="metric-card">

                <div class="metric-number">
                    {keyword_coverage:.1f}%
                </div>

                <div class="metric-label">
                    Job Keyword Coverage
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    with keyword_col2:

        st.markdown(
            f"""
            <div class="metric-card">

                <div class="metric-number">
                    {len(keyword_matched)}
                </div>

                <div class="metric-label">
                    Keywords Matched
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    with keyword_col3:

        st.markdown(
            f"""
            <div class="metric-card">

                <div class="metric-number">
                    {len(keyword_missing)}
                </div>

                <div class="metric-label">
                    Keywords Missing
                </div>

            </div>
            """,
            unsafe_allow_html=True,
        )


    if keyword_matched:

        st.markdown(
            "#### ✅ Matched Job Keywords"
        )

        st.markdown(
            render_skills(
                keyword_matched,
                "matched"
            ),
            unsafe_allow_html=True,
        )


    if keyword_missing:

        st.markdown(
            "#### ⚠️ Missing Job Keywords"
        )

        st.markdown(
            render_skills(
                keyword_missing,
                "missing"
            ),
            unsafe_allow_html=True,
        )


    # ========================================================
    # SKILL ANALYSIS
    # ========================================================

    st.divider()

    col1, col2 = st.columns(2)


    with col1:

        st.subheader(
            "✅ Matched Skills"
        )

        st.markdown(
            render_skills(
                matched_skills,
                "matched",
            ),
            unsafe_allow_html=True,
        )

        st.caption(
            f"{len(matched_skills)} required skill(s) "
            "found in the resume."
        )


    with col2:

        st.subheader(
            "❌ Missing Skills"
        )

        st.markdown(
            render_skills(
                missing_skills,
                "missing",
            ),
            unsafe_allow_html=True,
        )

        st.caption(
            f"{len(missing_skills)} required skill(s) "
            "were not detected in the resume."
        )


    # ========================================================
    # ALL RESUME SKILLS
    # ========================================================

    st.divider()

    st.subheader(
        "🧠 Skills Detected in Resume"
    )

    st.markdown(
        render_skills(
            resume_skills
        ),
        unsafe_allow_html=True,
    )

    st.caption(
        f"Total detected skills: "
        f"{len(resume_skills)}"
    )


    # ========================================================
    # REQUIRED SKILLS
    # ========================================================

    st.subheader(
        "🎯 Skills Detected in Job Description"
    )

    st.markdown(
        render_skills(
            required_skills
        ),
        unsafe_allow_html=True,
    )

    st.caption(
        f"Total required skills detected: "
        f"{len(required_skills)}"
    )


    # ========================================================
    # SUGGESTIONS
    # ========================================================

    st.divider()

    st.subheader(
        "💡 Resume Improvement Suggestions"
    )

    if suggestions:

        for suggestion in suggestions:

            st.markdown(
                f"""
                <div class="suggestion">
                    💡 {html.escape(suggestion)}
                </div>
                """,
                unsafe_allow_html=True,
            )

    else:

        st.success(
            "No additional suggestions were generated."
        )


    # ========================================================
    # SCORE EXPLANATION
    # ========================================================

    with st.expander(
        "ℹ️ How the score is calculated"
    ):

        st.write(
            "The analyzer combines two signals:"
        )

        st.write(
            "1. **Skill Match — 60% weight:** "
            "Required skills detected in the job "
            "description are compared with skills "
            "detected in the resume."
        )

        st.write(
            "2. **Text Similarity — 40% weight:** "
            "TF-IDF cosine similarity compares the "
            "resume and job-description text."
        )

        st.write(
            "Final Score = "
            "(Skill Match × 0.60) + "
            "(Text Similarity × 0.40)"
        )

        st.warning(
            "This score is an automated text-analysis "
            "estimate, not a hiring decision or guarantee."
        )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "AI Resume Analyzer • PDF / DOCX / TXT • "
    "Local analysis • No API key required"
)